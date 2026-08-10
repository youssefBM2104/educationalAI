import base64
import io
import logging
import re
import tempfile
from pathlib import Path

from backend.agents.state import LectureState
from backend.agents.slide_generation._common import build_image_index

logger = logging.getLogger(__name__)

_FORBIDDEN = r'<>:"/\|?*'


def _safe_stem(name: str, max_len: int = 120) -> str:
    cleaned = re.sub(f"[{re.escape(_FORBIDDEN)}]", "-", name or "")
    cleaned = re.sub(r"\s+", "_", cleaned.strip())
    cleaned = re.sub(r"[-_]{2,}", "_", cleaned).strip("._-")
    return (cleaned or "lecture")[:max_len]


def _slide_lines(slide: dict) -> list[str]:
    """Flatten a slide's content (bullets + layout-specific slots) into displayable text lines.
    A placeholder renderer until the visual .pptx-template step maps slots to real layout shapes."""
    lines = list(slide.get("bullets") or [])
    slots = slide.get("slots") or {}
    if slots.get("subtitle"):
        lines.append(str(slots["subtitle"]))
    if slots.get("stat"):
        lines.append(f"{slots['stat']} — {slots.get('explanation', '')}".strip(" —"))
    for card in slots.get("cards") or []:
        if isinstance(card, dict):
            lines.append(f"{card.get('label', '')}: {card.get('text', '')}".strip(": "))
        else:
            lines.append(str(card))
    if slots.get("caption"):
        lines.append(str(slots["caption"]))
    return lines


def _backfill_empty_slides(slides: list[dict], headings: list[dict]) -> int:
    """Safety net for planner drop-outs: the Slide Planner occasionally emits a slide with a title
    but no bullets/slots (its cards/bullets came back empty). Since the Composer already wrote real
    `content_points` per heading, backfill those as bullets so the slide is never blank. Deterministic
    — does not depend on the (non-deterministic) planner filling every slot."""
    by_heading: dict[str, list] = {}
    for h in headings or []:
        by_heading.setdefault(h.get("heading_id"), h.get("content_points") or [])
    fixed = 0
    for s in slides:
        if _slide_lines(s) or (s.get("image_refs") or []):
            continue                               # already has content (text or an image)
        points = by_heading.get(s.get("heading_id"))
        if points:
            s["bullets"] = list(points)
            fixed += 1
    return fixed


def _decode(image_index: dict, image_id: str) -> io.BytesIO | None:
    entry = image_index.get(image_id)
    if not entry or not entry.get("image_base64"):
        logger.warning("Export: image_id %s not found / no bytes — skipping", image_id)
        return None
    try:
        raw = entry["image_base64"]
        raw = raw.split(",", 1)[1] if raw.startswith("data:") else raw
        return io.BytesIO(base64.b64decode(raw))
    except Exception as e:
        logger.warning("Export: failed to decode image %s: %s", image_id, e)
        return None


def _open_pptx(template):
    """Open the per-request visual template (path or bytes) as the base deck, stripping its own
    demo slides but keeping its theme + layouts, so our content slides inherit the Canva look.
    Falls back to a blank deck when no template is given."""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    if not template:
        return Presentation()
    src = io.BytesIO(template) if isinstance(template, (bytes, bytearray)) else template
    prs = Presentation(src)

    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        rid = sld_id.get(qn("r:id"))
        if rid:
            try:
                prs.part.drop_rel(rid)
            except KeyError:
                pass
        sld_id_lst.remove(sld_id)
    return prs



_WANT = {
    "title":            (0, False, True),
    "title_bullets":    (1, False, False),
    "stat":             (1, False, False),
    "three_card":       (2, False, False),
    "two_column_image": (1, True,  False),
    "image_caption":    (0, True,  False),
}


def _roles(placeholders):
    """Split placeholders into (titles, bodies, pictures), ignoring chrome (date/footer/number)."""
    from pptx.enum.shapes import PP_PLACEHOLDER as PP
    titles, bodies, pics = [], [], []
    for p in placeholders:
        t = p.placeholder_format.type
        if t in (PP.TITLE, PP.CENTER_TITLE) or p.placeholder_format.idx == 0:
            titles.append(p)
        elif t == PP.PICTURE:
            pics.append(p)
        elif t in (PP.BODY, PP.OBJECT, PP.SUBTITLE):
            bodies.append(p)

    return titles, bodies, pics


def _catalog(prs):
    """Signature of every template layout: how many title/body/picture slots it offers."""
    from pptx.enum.shapes import PP_PLACEHOLDER as PP
    cat = []
    for lay in prs.slide_layouts:
        titles, bodies, pics = _roles(lay.placeholders)
        cat.append({
            "layout": lay,
            "n_title": len(titles),
            "n_body": len(bodies),
            "n_pic": len(pics),
            "center": any(p.placeholder_format.type == PP.CENTER_TITLE for p in titles),
        })
    return cat


def _choose(cat: list[dict], semantic: str, need_pic: bool = False):
    """Pick the template layout whose signature best fits a semantic slide type. `need_pic` is set by
    Export when the slide actually carries an image, forcing a layout that HAS a picture placeholder
    (so the image lands in a designed well instead of floating over the text). Pure scoring with a
    guaranteed fallback, so a template that lacks an ideal layout still renders (never breaks)."""
    want_min_body, want_pic, prefer_center = _WANT.get(semantic, (1, False, False))
    want_pic = want_pic or need_pic
    # If we truly need a picture well and the template has one, restrict to those layouts.
    candidates = cat
    if need_pic and any(c["n_pic"] >= 1 for c in cat if c["n_title"] > 0):
        candidates = [c for c in cat if c["n_pic"] >= 1]
    best, best_score = None, None
    for c in candidates:
        if c["n_title"] == 0:          # must be able to place the heading
            continue
        score = 0.0
        score += 3 if (want_pic and c["n_pic"] >= 1) else 0
        if not want_pic and c["n_pic"] >= 1:
            score -= 2                 # don't drop plain bullets onto a picture-well layout
        score += 2 if c["n_body"] >= want_min_body else -1.5 * (want_min_body - c["n_body"])
        score += 1 if (prefer_center and c["center"]) else 0
        score += 1 if (not prefer_center and not c["center"]) else 0
        score -= 0.1 * abs(c["n_body"] - max(want_min_body, 1))   # prefer the least over-provisioned
        if best_score is None or score > best_score:
            best, best_score = c, score
    return (best or candidates[0] if candidates else cat[0])["layout"]


def _card_line(card) -> str:
    if isinstance(card, dict):
        label = card.get("label") or card.get("title") or card.get("heading") or card.get("name") or ""
        text = card.get("text") or card.get("description") or card.get("body") or card.get("detail") or ""
        return f"{label}: {text}".strip(": ") or " ".join(str(v) for v in card.values() if v)
    return str(card)


def _write_lines(ph, lines: list) -> None:
    from pptx.enum.text import MSO_AUTO_SIZE
    tf = ph.text_frame
    tf.clear()
    tf.word_wrap = True
    try:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE   # shrink text to the template's box, no overflow
    except Exception:
        pass
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(line)


def _fill_bodies(slide_dict: dict, semantic: str, bodies: list) -> set:
    """Distribute content across the layout's body placeholders. Three parallel cards spread across
    parallel body slots when the layout offers them; everything else goes into the first body slot."""
    used = set()
    if not bodies:
        return used
    cards = (slide_dict.get("slots") or {}).get("cards") or []
    if semantic == "three_card" and cards and len(bodies) >= 2:
        for i, ph in enumerate(bodies):
            if i == len(bodies) - 1:
                chunk = cards[i:]                 # last slot absorbs any leftover cards
            else:
                chunk = [cards[i]] if i < len(cards) else []
            lines = [_card_line(c) for c in chunk]
            if not lines:
                continue                          # more slots than cards -> leave empty -> cleaned up
            _write_lines(ph, lines)
            used.add(ph.placeholder_format.idx)
        return used
    lines = _slide_lines(slide_dict)
    if lines:
        _write_lines(bodies[0], lines)
        used.add(bodies[0].placeholder_format.idx)
    return used


def _ph_box(slide, ph):
    """(left, top, width, height) of a placeholder, falling back to the matching LAYOUT placeholder
    when the slide copy inherits its geometry (its own left/top/width/height come back None)."""
    if None not in (ph.left, ph.top, ph.width, ph.height):
        return ph.left, ph.top, ph.width, ph.height
    idx = ph.placeholder_format.idx
    for lp in slide.slide_layout.placeholders:
        if lp.placeholder_format.idx == idx and None not in (lp.left, lp.top, lp.width, lp.height):
            return lp.left, lp.top, lp.width, lp.height
    return None


def _place_image(slide, slide_dict: dict, image_index: dict, pics: list) -> set:
    """Place the slide's image into the template's designed picture WELL: read the picture
    placeholder's box (from the layout if inherited), then add_picture fitted to that box, aspect
    preserved and centred — so it never floats over the title/body. `insert_picture` is avoided
    (it errors on some templates). Leaves the empty placeholder to be cleaned up afterwards."""
    from pptx.util import Inches
    used: set = set()
    iid = next((i for i in (slide_dict.get("image_refs") or [])), None)
    if not iid:
        return used
    buf = _decode(image_index, iid)
    if not buf:
        return used

    box = _ph_box(slide, pics[0]) if pics else None
    if box is None:                                   # template has no picture well -> safe corner
        box = (Inches(9.4), Inches(1.6), Inches(3.2), Inches(4.4))
    left, top, w, h = box

    pic = slide.shapes.add_picture(buf, left, top, width=w)
    if pic.height and h and pic.height > h:           # too tall for the well -> fit by height, keep aspect
        ar = pic.width / pic.height
        pic.height = int(h)
        pic.width = int(h * ar)
    pic.left = int(left + (w - pic.width) / 2)         # centre inside the well
    pic.top = int(top + (h - pic.height) / 2)
    return used                                        # picture placeholder stays unfilled -> removed


def _render_pptx(slides: list[dict], image_index: dict, out_path: Path, template=None) -> None:
    prs = _open_pptx(template)
    cat = _catalog(prs)

    for s in slides:
        semantic = s.get("layout") or "title_bullets"
        need_pic = any(_decode(image_index, iid) for iid in (s.get("image_refs") or []))
        slide = prs.slides.add_slide(_choose(cat, semantic, need_pic=need_pic))
        titles, bodies, pics = _roles(slide.placeholders)
        used = set()


        if titles:
            _write_lines(titles[0], [s.get("title", "")])
            used.add(titles[0].placeholder_format.idx)
        used |= _fill_bodies(s, semantic, bodies)
        used |= _place_image(slide, s, image_index, pics)


        for ph in list(slide.placeholders):
            if ph.placeholder_format.idx not in used:
                ph._element.getparent().remove(ph._element)

    prs.save(str(out_path))


def _render_pdf(slides: list[dict], image_index: dict, out_path: Path) -> None:
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader

    c = canvas.Canvas(str(out_path), pagesize=landscape(A4))
    w, h = landscape(A4)
    for s in slides:
        c.setFont("Helvetica-Bold", 22)
        c.drawString(0.6 * inch, h - 0.9 * inch, s.get("title", ""))
        c.setFont("Helvetica", 13)
        y = h - 1.5 * inch
        for b in _slide_lines(s):
            c.drawString(0.8 * inch, y, f"• {b}")
            y -= 0.35 * inch
        for iid in s.get("image_refs") or []:
            buf = _decode(image_index, iid)
            if buf:
                try:
                    c.drawImage(ImageReader(buf), w - 4.2 * inch, h - 4.5 * inch,
                                width=3.4 * inch, preserveAspectRatio=True, mask="auto")
                except Exception as e:
                    logger.warning("Export: pdf image %s failed: %s", iid, e)
                break
        c.showPage()
    c.save()


def export_agent(state: LectureState) -> dict:
    passed = state.get("verification_passed", False)
    if passed:
        deck = state.get("lecture_slides") or {}
        composer = state.get("composer_output") or {}
        status = "verified"
    else:
        best = state.get("best_attempt") or {}
        deck = best.get("lecture_slides") or state.get("lecture_slides") or {}
        composer = best.get("composer_output") or state.get("composer_output") or {}
        status = "best_attempt_unverified"

    slides = deck.get("slides", [])
    filled = _backfill_empty_slides(slides, composer.get("headings", []))
    if filled:
        logger.info("Export: backfilled %d empty slide(s) from composer content_points", filled)
    image_index = build_image_index(state.get("rag_chunks"))
    fmt = state.get("output_format", "pptx")

    stem = _safe_stem(state.get("query", "lecture"))
    out_dir = Path(tempfile.gettempdir()) / "slide_export"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{stem}.{fmt}"

    if fmt == "pdf":
        _render_pdf(slides, image_index, out_path)   # template applies to pptx only
    else:
        _render_pptx(slides, image_index, out_path, template=state.get("slide_template_pptx"))

    data = out_path.read_bytes()
    logger.info("Export: %d slides -> %s (%s, %d bytes)", len(slides), out_path, status, len(data))
    return {
        "lecture_output_path": str(out_path),
        "lecture_output_bytes": data,
        "export_status": status,
        "final_output": f"Slide deck ({status}): {out_path}",
    }
