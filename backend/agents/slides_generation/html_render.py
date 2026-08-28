from __future__ import annotations

import logging
import subprocess
import tempfile
from pathlib import Path

logger = logging.getLogger(__name__)

WIDTH, HEIGHT = 1280, 720   # 16:9

_BROWSERS = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
]


def _find_browser() -> str:
    for p in _BROWSERS:
        if Path(p).exists():
            return p
    raise RuntimeError("No Chrome/Edge found for HTML rendering. Install Chrome or Edge.")


def html_to_png(html: str, out_png: str | Path, width: int = WIDTH, height: int = HEIGHT) -> Path:
    """Rasterise one HTML string to a PNG via headless Chromium screenshot."""
    out_png = Path(out_png)
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False, encoding="utf-8") as f:
        f.write(html)
        html_path = Path(f.name)
    try:
        cmd = [
            _find_browser(),
            "--headless=new", "--disable-gpu", "--hide-scrollbars",
            "--force-device-scale-factor=1",
            f"--window-size={width},{height}",
            f"--screenshot={out_png}",
            html_path.as_uri(),
        ]
        proc = subprocess.run(cmd, capture_output=True, timeout=60)
        if not out_png.exists():
            raise RuntimeError(
                f"headless screenshot produced no file (exit {proc.returncode}): "
                f"{proc.stderr.decode(errors='ignore')[:300]}"
            )
        return out_png
    finally:
        html_path.unlink(missing_ok=True)


def pngs_to_pptx(png_paths: list[str | Path], out_pptx: str | Path) -> Path:
    """Assemble PNG slides into a 16:9 .pptx, one full-bleed image per slide."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for png in png_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(str(out_pptx))
    return Path(out_pptx)
