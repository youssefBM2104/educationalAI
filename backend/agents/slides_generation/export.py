import io
import logging
import os
import re
import tempfile
from pathlib import Path

from backend.agents.state import LectureState

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Filename sanitization
# ---------------------------------------------------------------------------
# Characters forbidden (or dangerous) in Windows filenames: < > : " / \ | ? *
# ':' is especially dangerous: NTFS silently interprets "Foo: Bar.pptx" as
# "create a file named 'Foo' with an Alternate Data Stream named
# ' Bar.pptx'". That produces exactly the symptom reported — a visible
# 0-byte file called "Thread_Synchronization" with no extension, while the
# real bytes get written into a hidden stream nobody can see in Explorer.
_WINDOWS_FORBIDDEN_CHARS = r'<>:"/\|?*'
_WINDOWS_RESERVED_NAMES = {
    "CON", "PRN", "AUX", "NUL",
    *(f"COM{i}" for i in range(1, 10)),
    *(f"LPT{i}" for i in range(1, 10)),
}


def _sanitize_filename(name: str, max_len: int = 150) -> str:
    """Turn an arbitrary course title into a safe, cross-platform filename stem."""
    # Replace forbidden characters with '-' (keeps things readable)
    cleaned = re.sub(f"[{re.escape(_WINDOWS_FORBIDDEN_CHARS)}]", "-", name)
    # Collapse whitespace to single underscores
    cleaned = re.sub(r"\s+", "_", cleaned.strip())
    # Drop any remaining control characters
    cleaned = re.sub(r"[\x00-\x1f]", "", cleaned)
    # Collapse repeated separators produced by the substitutions above
    cleaned = re.sub(r"[-_]{2,}", "_", cleaned).strip("._-")

    if not cleaned:
        cleaned = "lecture"

    if cleaned.upper() in _WINDOWS_RESERVED_NAMES:
        cleaned = f"_{cleaned}"

    return cleaned[:max_len]

# ---------------------------------------------------------------------------
# Optional heavy imports — only fail at call time, not at import time
# ---------------------------------------------------------------------------
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    _PPTX_AVAILABLE = True
except ImportError:
    _PPTX_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
    from reportlab.lib.enums import TA_LEFT
    _REPORTLAB_AVAILABLE = True
except ImportError:
    _REPORTLAB_AVAILABLE = False


# ---------------------------------------------------------------------------
# PPTX builder
# ---------------------------------------------------------------------------

def _build_pptx(slides: list[dict], course_title: str) -> bytes:
    if not _PPTX_AVAILABLE:
        raise ImportError("python-pptx is not installed. Run: pip install python-pptx")

    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    blank_layout  = prs.slide_layouts[6]   # completely blank
    title_layout  = prs.slide_layouts[0]   # title + subtitle
    body_layout   = prs.slide_layouts[1]   # title + content

    for slide_data in slides:
        stype   = slide_data.get("type", "content")
        title   = slide_data.get("title", "")
        bullets = slide_data.get("bullets", [])
        notes   = slide_data.get("speaker_notes", "")
        hint    = slide_data.get("visual_hint", "")

        # --- Title slide ---
        if stype == "title":
            slide = prs.slides.add_slide(title_layout)
            slide.shapes.title.text = course_title
            slide.placeholders[1].text = title if title != course_title else ""

        # --- Content / Summary slide ---
        else:
            slide = prs.slides.add_slide(body_layout)
            slide.shapes.title.text = title

            tf = slide.placeholders[1].text_frame
            tf.clear()

            for i, bullet in enumerate(bullets):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text  = bullet
                para.level = 0
                para.font.size = Pt(18)

            # Visual hint as a smaller italic paragraph at the bottom
            if hint:
                p = tf.add_paragraph()
                p.text = f"💡 {hint}"
                p.font.size   = Pt(12)
                p.font.italic = True
                p.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        # Speaker notes
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PDF builder
# ---------------------------------------------------------------------------

def _build_pdf(slides: list[dict], course_title: str) -> bytes:
    if not _REPORTLAB_AVAILABLE:
        raise ImportError("reportlab is not installed. Run: pip install reportlab")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=2*cm, rightMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)

    styles  = getSampleStyleSheet()
    h1      = styles["Heading1"]
    h2      = styles["Heading2"]
    normal  = styles["Normal"]
    italic  = ParagraphStyle("italic", parent=normal, fontName="Helvetica-Oblique",
                             textColor=(0.5, 0.5, 0.5), fontSize=9)

    story = []

    for slide_data in slides:
        stype   = slide_data.get("type", "content")
        title   = slide_data.get("title", "")
        bullets = slide_data.get("bullets", [])
        notes   = slide_data.get("speaker_notes", "")
        hint    = slide_data.get("visual_hint", "")

        if stype == "title":
            story.append(Paragraph(course_title, h1))
            if title and title != course_title:
                story.append(Paragraph(title, h2))
        else:
            story.append(Paragraph(title, h2))

        if bullets:
            items = [ListItem(Paragraph(b, normal), bulletColor=(0, 0, 0)) for b in bullets]
            story.append(ListFlowable(items, bulletType="bullet"))

        if hint:
            story.append(Spacer(1, 0.3*cm))
            story.append(Paragraph(f"💡 {hint}", italic))

        # Speaker notes as footnote-style block
        if notes:
            story.append(Spacer(1, 0.2*cm))
            story.append(Paragraph(f"<i>Note: {notes}</i>", italic))

        story.append(Spacer(1, 0.8*cm))

    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Node
# ---------------------------------------------------------------------------

def export(state: LectureState) -> dict:
    deck = state["lecture_slides"]
    slides = deck["slides"]
    course_title = state["lecture_plan"]["course_title"]
    fmt = state.get("output_format", "pptx")

    safe_stem = _sanitize_filename(course_title)

    if fmt == "pdf":
        raw   = _build_pdf(slides, course_title)
        fname = f"{safe_stem}.pdf"
    else:
        raw   = _build_pptx(slides, course_title)
        fname = f"{safe_stem}.pptx"

    # --- Sanity check: never silently write an empty/corrupt file ---
    if not raw:
        raise RuntimeError(
            f"Export produced 0 bytes for format={fmt!r} — the presentation/document "
            "was empty or the builder failed silently."
        )

    # --- Resolve output directory (absolute path avoids surprises tied to CWD) ---
    output_dir = Path("outputs/lectures").resolve()
    try:
        output_dir.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        raise RuntimeError(f"Could not create output directory {output_dir}: {e}") from e

    output_path = output_dir / fname

    # --- Atomic write: write to a temp file in the same directory, then
    #     replace the target. This avoids partially-written files if the
    #     process is interrupted, and avoids permission surprises since the
    #     temp file is created explicitly rather than truncate-on-open. ---
    try:
        with tempfile.NamedTemporaryFile(
            dir=output_dir, prefix=f".{safe_stem}_", suffix=".tmp", delete=False
        ) as tmp_f:
            tmp_f.write(raw)
            tmp_f.flush()
            os.fsync(tmp_f.fileno())
            tmp_path = Path(tmp_f.name)

        os.replace(tmp_path, output_path)  # atomic on both Windows and POSIX
    except OSError as e:
        # Clean up the temp file if the replace step failed
        try:
            if "tmp_path" in locals() and tmp_path.exists():
                tmp_path.unlink()
        except OSError:
            pass
        raise RuntimeError(
            f"Failed to write export file to {output_path} "
            f"(check disk space and write permissions on {output_dir}): {e}"
        ) from e

    # --- Verify what actually landed on disk matches what we intended ---
    actual_size = output_path.stat().st_size
    if actual_size != len(raw):
        raise RuntimeError(
            f"Export size mismatch for {output_path}: expected {len(raw)} bytes, "
            f"found {actual_size} bytes on disk."
        )

    logger.info(
        "Export: format=%s slides=%d path=%s size=%d bytes",
        fmt, len(slides), output_path, actual_size,
    )

    return {
        "lecture_output_path":  str(output_path),
        "lecture_output_bytes": raw,   # available for direct API streaming if needed
    }